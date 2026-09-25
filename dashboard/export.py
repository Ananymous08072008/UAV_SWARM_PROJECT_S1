"""
dashboard/export.py
Packs one simulation run into a single Excel workbook for download.

The mission studio writes nothing to disk - a shared server would litter it with
a directory per visitor, and a hosted one loses the disk anyway - so the
workbook is built in memory and handed to the browser.

    Mission metrics    the grouped summary used in the report, plus run details
    Metrics over time  one row per metric sample
    Event log          every event of the run
"""

from __future__ import annotations

import io
import json
import re
from datetime import datetime, timezone
from typing import Any, Mapping, Optional

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

EVENT_COLUMNS = (("Seq", 8), ("Time (s)", 10), ("Type", 22), ("Severity", 10), ("UAV", 6),
                 ("PoI", 12), ("Message", 80), ("Details", 60))

# Anything outside this is stripped from the download name. Scenario names reach
# us from YAML and from the browser, and the name goes into a Content-Disposition
# header - an unescaped quote or newline there is a header injection.
_UNSAFE = re.compile(r"[^A-Za-z0-9._-]+")
_BOLD = Font(bold=True)


def safe_name(text: str, fallback: str = "run") -> str:
    cleaned = _UNSAFE.sub("-", str(text)).strip("-._")
    return cleaned[:60] or fallback


def _label(key: str) -> str:
    return key.replace("_", " ").capitalize()


def _cell(ws, value: Any, bold: bool = False) -> WriteOnlyCell:
    if isinstance(value, (Mapping, list, tuple)):
        value = json.dumps(value, sort_keys=True, default=str)
    elif value is not None and not isinstance(value, (bool, int, float)):
        value = str(value)
    cell = WriteOnlyCell(ws, value=value)
    if isinstance(value, str):
        # Event messages and PoI names can come from the browser. A string that
        # starts with "=" would otherwise be stored as a live formula.
        cell.data_type = "s"
    if bold:
        cell.font = _BOLD
    return cell


def _row(ws, values, bold: bool = False) -> None:
    ws.append([_cell(ws, v, bold) for v in values])


def _sheet(wb: Workbook, title: str, widths: list[float]):
    ws = wb.create_sheet(title)
    for i, width in enumerate(widths):
        ws.column_dimensions[get_column_letter(i + 1)].width = width
    return ws


def run_stem(sim, session_id: Optional[str] = None) -> str:
    """The workbook's base name: time, scenario, mode (and session)."""
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    parts = [stamp, safe_name(sim.world.scenario.name, "scenario"), safe_name(sim.mode, "mode")]
    if session_id:
        parts.append(safe_name(session_id, "session"))
    return "_".join(parts)


def build_workbook(sim, session_id: Optional[str] = None,
                   extra: Optional[Mapping[str, Any]] = None) -> tuple[str, bytes]:
    """
    Build the workbook for ``sim`` and return ``(filename, bytes)``.

    Safe to call on a run that is still going: every list is snapshotted first,
    because the simulation thread owns them and keeps appending.
    """
    stem = run_stem(sim, session_id)
    rows = list(sim.metrics.rows())

    # A finished run already computed its summary; a running one is aggregated now.
    # A crashed run still has an event log worth downloading, so a summary that
    # cannot be produced must not take the download down with it.
    summary, summary_error = sim.summary, None
    if summary is None:
        try:
            summary = sim.metrics.summary()
        except Exception as exc:
            summary, summary_error = {}, f"{type(exc).__name__}: {exc}"

    logger = getattr(sim, "logger", None)
    if logger is not None:
        records = list(logger.records)
        log_note = (f"incomplete - the oldest {logger.dropped} events were dropped"
                    if logger.dropped else "complete")
    else:
        # No recorder attached: the bus keeps only its most recent events, so say
        # so rather than presenting the tail as the whole run.
        records = [e.to_dict() for e in sim.world.events.history()]
        log_note = "may be incomplete - only the most recent events were kept"

    wb = Workbook(write_only=True)

    ws = _sheet(wb, "Mission metrics", [16, 30, 24])
    ws.freeze_panes = "A2"
    _row(ws, ("Category", "Metric", "Value"), bold=True)
    details = {"exported_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
               "run_finished": sim.summary is not None, "sim_time_s": round(sim.world.t, 2),
               "event_log": log_note}
    if session_id:
        details["session_id"] = session_id
    details.update(extra or {})
    if summary_error is not None:
        details["summary_error"] = summary_error
    for key, value in details.items():
        _row(ws, ("Export", _label(key), value))
    for group, values in summary.items():
        for key, value in values.items():
            _row(ws, (_label(group), _label(key), value))

    columns = list(rows[0]) if rows else []
    ws = _sheet(wb, "Metrics over time", [14] * len(columns))
    ws.freeze_panes = "A2"
    _row(ws, columns, bold=True)
    for sample in rows:
        _row(ws, [sample.get(c) for c in columns])

    ws = _sheet(wb, "Event log", [w for _, w in EVENT_COLUMNS])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{len(records) + 1}"
    _row(ws, [name for name, _ in EVENT_COLUMNS], bold=True)
    for r in records:
        _row(ws, (r["seq"], r["t_s"], r["type"], r["severity"], r["uav_id"], r["poi_id"],
                  r["message"], r.get("data") or None))

    buf = io.BytesIO()
    wb.save(buf)
    return f"{stem}.xlsx", buf.getvalue()
