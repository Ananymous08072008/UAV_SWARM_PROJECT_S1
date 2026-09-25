"""Run download: the mission metrics and event log of a run, as an Excel workbook.

The point of the feature is that the workbook is *complete*. Both the EventBus
and the LiveHub are bounded ring buffers, so a run of any length holds only its
tail in either. These tests exist mostly to prove the download does not inherit
that truncation, and says so when it cannot avoid it.
"""

from __future__ import annotations

import io
import time

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from core.events import EventBus, EventType
from dashboard.api import create_app
from dashboard.export import XLSX_MEDIA_TYPE, build_workbook, safe_name
from dashboard.session import SessionManager
from dashboard.websocket import LiveHub
from swarm_logging.event_logger import EventLogger
from tests.helpers import make_sim, run_until

SHEETS = ["Mission metrics", "Metrics over time", "Event log"]

SPEC = {"uav_count": 4, "duration_s": 60.0, "speed": 20.0,
        "pois": [{"x_m": 400.0, "y_m": 200.0, "priority": 4, "survey_time_s": 20.0}]}


def open_book(blob: bytes):
    return load_workbook(io.BytesIO(blob), read_only=True)


def table(book, sheet: str) -> list[tuple]:
    return [tuple(row) for row in book[sheet].iter_rows(values_only=True)]


def metrics(book) -> dict[tuple[str, str], object]:
    """Mission metrics as {(category, metric): value}. An empty value reads back as a short row."""
    out = {}
    for row in table(book, "Mission metrics")[1:]:
        category, metric, value = (row + (None, None, None))[:3]
        out[(category, metric)] = value
    return out


def wait_until(predicate, timeout_s: float = 10.0) -> bool:
    deadline = time.perf_counter() + timeout_s
    while time.perf_counter() < deadline:
        if predicate():
            return True
        time.sleep(0.05)
    return False


@pytest.fixture()
def finished_run():
    hub = LiveHub()
    sim = make_sim(hub=hub)
    sim.start()
    run_until(sim, 40)
    sim.finish("duration reached")
    return sim, hub


# ------------------------------------------------------------------ the workbook
def test_the_workbook_holds_only_the_metrics_and_the_event_log(finished_run):
    sim, _ = finished_run
    name, blob = build_workbook(sim)
    assert name.endswith(".xlsx")
    assert open_book(blob).sheetnames == SHEETS


def test_mission_metrics_carry_every_evaluation_section(finished_run):
    sim, _ = finished_run
    rows = metrics(open_book(build_workbook(sim)[1]))
    categories = {c for c, _ in rows}
    assert {"Run", "Mission", "Communication", "Resilience", "Safety", "Efficiency"} <= categories
    assert rows[("Mission", "Completion rate")] == sim.summary["mission"]["completion_rate"]
    assert rows[("Run", "Seed")] == sim.world.seed        # numbers stay numbers, not text


def test_metrics_over_time_has_one_row_per_sample(finished_run):
    sim, _ = finished_run
    rows = table(open_book(build_workbook(sim)[1]), "Metrics over time")
    assert len(rows) - 1 == len(sim.metrics.rows())
    assert {"t_s", "connectivity_ratio", "mean_route_pdr", "mean_latency_ms"} <= set(rows[0])


def test_the_event_log_lists_every_event_in_order(finished_run):
    sim, _ = finished_run
    rows = table(open_book(build_workbook(sim)[1]), "Event log")
    assert rows[0] == ("Seq", "Time (s)", "Type", "Severity", "UAV", "PoI", "Message", "Details")
    seqs = [r[0] for r in rows[1:]]
    assert len(seqs) == len(sim.logger.records) and seqs == sorted(seqs)
    assert rows[1][2] == "SIM_STARTED"


def test_text_that_looks_like_a_formula_stays_text():
    # PoI names and messages can come from the browser; "=..." must not become a formula.
    sim = make_sim(hub=LiveHub())
    sim.start()
    sim.world.inject("add_poi", {"id": "=HYPERLINK(\"http://x\")", "position_m": [100, 100]})
    book = load_workbook(io.BytesIO(build_workbook(sim)[1]))    # full mode exposes data_type
    poi_cells = [row[5] for row in book["Event log"].iter_rows(min_row=2) if row[5].value]
    assert poi_cells and all(c.data_type == "s" for c in poi_cells)


# --------------------------------------------------------------- completeness
def test_the_workbook_keeps_events_the_live_hub_has_already_dropped():
    # The hub is a ring buffer. Exporting from it would silently hand over the
    # tail of a run and call it the run.
    hub = LiveHub(max_events=10)
    sim = make_sim(hub=hub)
    sim.start()
    run_until(sim, 40)
    sim.push(force=True)
    sim.finish("duration reached")

    book = open_book(build_workbook(sim)[1])
    events = table(book, "Event log")[1:]
    assert len(hub.events_since(0)) == 10            # the hub kept only its last ten
    assert len(events) > 10                          # the workbook kept all of them
    assert events[0][0] == 1                         # including the very first
    assert metrics(book)[("Export", "Event log")] == "complete"


def test_a_truncated_log_announces_itself_instead_of_lying():
    bus = EventBus()
    logger = EventLogger(None, bus, max_records=5)
    for i in range(20):
        bus.publish(float(i), EventType.UAV_ARRIVED, f"event {i}")
    assert len(logger.records) == 5
    assert logger.dropped == 15
    assert [r["message"] for r in logger.records][0] == "event 15"   # the most recent


def test_a_run_with_no_recorder_says_its_log_may_be_incomplete():
    # Constructed without a hub, so nothing recorded it; the bus is all there is.
    sim = make_sim()
    sim.start()
    run_until(sim, 10)
    note = metrics(open_book(build_workbook(sim)[1]))[("Export", "Event log")]
    assert "incomplete" in note


# -------------------------------------------------------------- the event logger
def test_a_logger_without_a_directory_writes_nothing(tmp_path):
    bus = EventBus()
    logger = EventLogger(None, bus)
    bus.publish(1.0, EventType.UAV_ARRIVED, "hello")
    logger.close()
    assert len(logger.records) == 1
    assert list(tmp_path.iterdir()) == []


def test_a_logger_with_a_directory_still_writes_both_files(tmp_path):
    bus = EventBus()
    logger = EventLogger(tmp_path / "run", bus)
    bus.publish(1.0, EventType.UAV_ARRIVED, "hello")
    logger.close()
    assert (tmp_path / "run" / "events.jsonl").read_text().strip()
    assert (tmp_path / "run" / "events.csv").read_text().strip()
    assert logger.dropped == 0


# ------------------------------------------------------- single-simulation route
def test_the_dashboard_serves_its_run_as_a_download(finished_run):
    sim, hub = finished_run
    with TestClient(create_app(hub)) as client:
        response = client.get("/api/export")
    assert response.status_code == 200
    assert response.headers["content-type"] == XLSX_MEDIA_TYPE
    assert "attachment" in response.headers["content-disposition"]
    assert open_book(response.content).sheetnames == SHEETS


def test_exporting_with_no_simulation_attached_is_a_conflict():
    with TestClient(create_app(LiveHub())) as client:
        assert client.get("/api/export").status_code == 409


def test_a_simulation_registers_itself_with_the_hub_it_is_given():
    hub = LiveHub()
    assert hub.simulation is None
    sim = make_sim(hub=hub)
    assert hub.simulation is sim
    assert sim.logger is not None       # and starts recording for the download


# ------------------------------------------------------------------ studio route
@pytest.fixture()
def manager():
    mgr = SessionManager(max_sessions=3)
    yield mgr
    mgr.shutdown()


@pytest.fixture()
def client(manager):
    with TestClient(create_app(manager=manager)) as test_client:
        yield test_client


def test_a_session_run_can_be_downloaded(client):
    session_id = client.post("/api/sessions", json=SPEC).json()["id"]
    assert wait_until(lambda: client.get(f"/api/sessions/{session_id}").json()["state"] == "finished",
                      timeout_s=30.0)
    response = client.get(f"/api/sessions/{session_id}/export")
    assert response.status_code == 200
    book = open_book(response.content)
    rows = metrics(book)
    assert rows[("Export", "Session id")] == session_id
    assert rows[("Export", "Run finished")] is True
    assert len(table(book, "Metrics over time")) > 1


def test_a_run_still_going_can_be_downloaded(client):
    session_id = client.post("/api/sessions", json=SPEC).json()["id"]
    assert wait_until(lambda: client.get(f"/api/sessions/{session_id}").json()["t_s"] > 2)
    rows = metrics(open_book(client.get(f"/api/sessions/{session_id}/export").content))
    assert rows[("Export", "Run finished")] is False
    assert rows[("Export", "State")] in ("running", "paused", "finished")


def test_downloading_an_unknown_session_is_404(client):
    assert client.get("/api/sessions/nosuchthing/export").status_code == 404


# ---------------------------------------------------------------------- the name
@pytest.mark.parametrize("hostile", [
    'evil" filename',
    "evil\r\nX-Injected: yes",
    "../../etc/passwd",
    "nåme with spaces",
])
def test_a_scenario_name_cannot_break_out_of_the_download_header(hostile):
    cleaned = safe_name(hostile)
    assert not set(cleaned) & set('"\r\n/\\ ')
    assert cleaned


def test_an_empty_name_still_produces_a_usable_filename():
    assert safe_name("", fallback="run") == "run"
    assert safe_name("!!!", fallback="run") == "run"


def test_the_download_filename_is_built_from_the_run(finished_run):
    sim, hub = finished_run
    with TestClient(create_app(hub)) as client:
        disposition = client.get("/api/export").headers["content-disposition"]
    assert disposition.count('"') == 2               # exactly the quotes we put there
    assert sim.mode in disposition
    assert disposition.endswith('.xlsx"')
